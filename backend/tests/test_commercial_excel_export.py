from __future__ import annotations

from datetime import datetime, timedelta, timezone
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import Session

from app.database import Base, get_db
from app.main import app
from app.models import (
    CommercialRelationship,
    ContactEvidence,
    ContactPoint,
    ObservedJobOffer,
    PersonContact,
)
from app.services.search_runs import SearchRunOrchestrator


NOW = datetime(2026, 9, 23, 10, tzinfo=timezone.utc)


@pytest.fixture
def session(tmp_path):
    engine = create_engine(
        f"sqlite:///{tmp_path / 'commercial-export.sqlite3'}",
        connect_args={"check_same_thread": False},
    )
    Base.metadata.create_all(engine)
    with Session(engine) as database_session:
        yield database_session
    engine.dispose()


@pytest.fixture
def client(session):
    def override_get_db():
        yield session

    app.dependency_overrides[get_db] = override_get_db
    with TestClient(app) as test_client:
        yield test_client
    app.dependency_overrides.clear()


def add_offer(
    session: Session,
    identifier: str,
    company: str,
    *,
    source: str = "france_travail",
    title: str = "Technicien",
    first_seen_hours_ago: int = 2,
    created_days_ago: int = 3,
    url: str | None = None,
    contract_type: str | None = None,
    salary: str | None = None,
    observation_count: int = 1,
) -> None:
    first_seen = NOW - timedelta(hours=first_seen_hours_ago)
    session.add(ObservedJobOffer(
        source=source,
        source_offer_id=identifier,
        title=title,
        company_name=company,
        location_label="Créteil",
        commune="Créteil",
        department_code="94",
        created_at=(NOW - timedelta(days=created_days_ago)).isoformat().replace("+00:00", "Z"),
        contract_type=contract_type,
        salary=salary,
        source_url=url,
        first_seen_at=first_seen,
        last_seen_at=NOW,
        last_changed_at=first_seen,
        is_active=True,
        observation_count=observation_count,
    ))
    session.flush()


def add_contact_facts(session: Session, company_key: str, company_name: str) -> None:
    person = PersonContact(
        company_key=company_key,
        scope="company",
        local_key=None,
        siren=None,
        organization_name_snapshot=company_name,
        local_commune_snapshot=None,
        local_location_label_snapshot=None,
        full_name="Élodie Martin",
        normalized_name="elodie martin",
        job_title="Responsable RH",
        relevance_role="hr",
        confidence_level="confirmed",
        verification_status="source_verified",
        attribution_reason="Fonction publiée sur le site officiel.",
        fingerprint=f"person-{company_key}",
        first_observed_at=NOW,
        last_observed_at=NOW,
        is_active=True,
    )
    session.add(person)
    session.flush()
    values = (
        ("email", "rh@societe-elegante.test"),
        ("phone", "+33102030405"),
        ("professional_url", "javascript:alert(1)"),
    )
    for index, (contact_type, value) in enumerate(values):
        point = ContactPoint(
            company_key=company_key,
            scope="company",
            local_key=None,
            siren=None,
            organization_name_snapshot=company_name,
            local_commune_snapshot=None,
            local_location_label_snapshot=None,
            contact_type=contact_type,
            value=value,
            normalized_value=value,
            confidence_level="confirmed",
            verification_status="source_verified",
            attribution_reason="Coordonnée professionnelle publique.",
            person_contact_id=person.id,
            fingerprint=f"point-{company_key}-{index}",
            first_observed_at=NOW,
            last_observed_at=NOW,
            is_active=True,
        )
        session.add(point)
        session.flush()
        session.add(ContactEvidence(
            contact_point_id=point.id,
            person_contact_id=None,
            provider="official_web",
            source_name="site_officiel",
            source_url=f"https://societe-elegante.test/contact#{index}",
            source_identifier=None,
            observed_at=NOW,
            evidence_reason="public contact",
            excerpt=None,
            fingerprint=f"evidence-{company_key}-{index}",
        ))
    session.commit()


def workbook_from(response):
    assert response.status_code == 200, response.text
    assert response.content[:2] == b"PK"
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert response.headers["content-disposition"].endswith('.xlsx"')
    return load_workbook(BytesIO(response.content))


def sheet_records(sheet):
    headers = [cell.value for cell in sheet[1]]
    return [dict(zip(headers, row)) for row in sheet.iter_rows(min_row=2, values_only=True)]


def test_endpoint_returns_safe_openable_structured_xlsx(client, session):
    company = "Société Élégante"
    add_offer(
        session, "ft-1", company, source="france_travail", title="Auxiliaire de vie H/F",
        url="https://francetravail.test/offre/1", contract_type="CDI", salary="35 000 €",
        observation_count=3,
    )
    add_offer(
        session, "linkedin-1", company, source="linkedin", title="Auxiliaire de vie",
        url="https://linkedin.test/jobs/1", observation_count=2,
    )
    add_offer(
        session, "ft-2", company, source="france_travail", title="Responsable RH",
        url="https://francetravail.test/offre/2",
    )
    add_offer(session, "missing", "Sans Contact", url=None)
    add_offer(session, "formula", "=2+2", url="https://source.test/formula")
    add_contact_facts(session, "société élégante", company)
    session.add(CommercialRelationship(
        company_key="société élégante",
        siren=None,
        company_name_snapshot=company,
        status="contacted",
        last_contact_at=NOW - timedelta(days=1),
        next_action_at=NOW + timedelta(days=2),
        is_active=True,
        created_at=NOW,
        updated_at=NOW,
    ))
    session.commit()

    workbook = workbook_from(client.get(
        "/api/v1/commercial-leads/export.xlsx?include_excluded=true"
    ))

    assert workbook.sheetnames == ["Prospects", "Offres", "Contacts"]
    assert [cell.value for cell in workbook["Prospects"][1]] == [
        "Entreprise", "SIREN", "Site web", "Secteur", "Localisation principale",
        "Taille", "Score /100", "Catégorie de score", "Nombre d’offres actives",
        "Dernière nouveauté", "Type de nouveauté", "Offre publiée la plus récente", "Rôles recrutés",
        "Relation employeur", "Statut commercial", "Dernier contact", "Prochaine relance",
        "Contact prioritaire", "Fonction du contact", "Canal recommandé", "Email",
        "Téléphone", "URL professionnelle", "Confiance contact",
        "Pourquoi ce prospect est intéressant", "Warnings / limites",
        "Données manquantes utiles", "Sources / provenance", "URLs de preuve",
    ]
    assert [cell.value for cell in workbook["Offres"][1]] == [
        "Entreprise", "SIREN", "Intitulé", "Localisation", "Commune", "Contrat",
        "Salaire", "Date de publication", "Première détection locale", "Dernière observation",
        "Âge (jours)", "Nombre d’observations", "Nombre de sources", "Sources",
        "URLs de preuve", "Nouvelle récemment", "Identifiants source",
    ]
    assert [cell.value for cell in workbook["Contacts"][1]] == [
        "Entreprise", "SIREN", "Nom / personne", "Fonction", "Pertinence", "Email",
        "Téléphone", "URL professionnelle", "Type de canal", "Scope", "Confiance",
        "Vérification", "Source / provenance", "URL de preuve", "Contexte utile",
        "Incertitude / warning",
    ]
    prospects = sheet_records(workbook["Prospects"])
    assert sum(row["Entreprise"] == company for row in prospects) == 1
    elegant = next(row for row in prospects if row["Entreprise"] == company)
    missing = next(row for row in prospects if row["Entreprise"] == "Sans Contact")
    injected = next(row for row in prospects if row["Entreprise"] == "'=2+2")
    assert elegant["Statut commercial"] == "Contacté"
    assert elegant["Email"] == "rh@societe-elegante.test"
    assert elegant["Téléphone"] == "+33102030405"
    assert elegant["Contact prioritaire"] == "Élodie Martin"
    assert elegant["Score /100"] == int(elegant["Score /100"])
    assert isinstance(elegant["Dernier contact"], datetime)
    assert missing["SIREN"] is None and missing["Email"] is None
    assert injected["Entreprise"].startswith("'")
    injected_row = next(
        row for row in range(2, workbook["Prospects"].max_row + 1)
        if workbook["Prospects"].cell(row, 1).value == "'=2+2"
    )
    assert workbook["Prospects"].cell(injected_row, 1).data_type == "s"

    offers = [row for row in sheet_records(workbook["Offres"]) if row["Entreprise"] == company]
    assert len(offers) == 2
    multi = next(row for row in offers if row["Intitulé"] == "Auxiliaire de vie H/F")
    assert multi["Nombre de sources"] == 2
    assert multi["Nombre d’observations"] == 5
    assert "france_travail" in multi["Sources"] and "linkedin" in multi["Sources"]
    assert "https://francetravail.test/offre/1" in multi["URLs de preuve"]
    assert "https://linkedin.test/jobs/1" in multi["URLs de preuve"]
    assert isinstance(multi["Date de publication"], datetime)
    assert isinstance(multi["Première détection locale"], datetime)
    assert isinstance(multi["Âge (jours)"], int)

    contacts = sheet_records(workbook["Contacts"])
    email = next(row for row in contacts if row["Email"] == "rh@societe-elegante.test")
    phone = next(row for row in contacts if row["Téléphone"] == "+33102030405")
    unsafe_url_row_index = next(
        index for index, row in enumerate(contacts, start=2)
        if row["URL professionnelle"] == "javascript:alert(1)"
    )
    unsafe_url_column = [cell.value for cell in workbook["Contacts"][1]].index("URL professionnelle") + 1
    proof_url_column = [cell.value for cell in workbook["Contacts"][1]].index("URL de preuve") + 1
    email_row_index = next(
        index for index, row in enumerate(contacts, start=2)
        if row["Email"] == "rh@societe-elegante.test"
    )
    assert email["Nom / personne"] == "Élodie Martin"
    assert phone["Fonction"] == "Responsable RH"
    assert email["URL de preuve"].startswith("https://")
    assert workbook["Contacts"].cell(email_row_index, proof_url_column).hyperlink.target.startswith("https://")
    assert workbook["Contacts"].cell(unsafe_url_row_index, unsafe_url_column).hyperlink is None
    assert all(row["Email"] is None for row in contacts if row["Entreprise"] == "Sans Contact")

    for sheet in workbook.worksheets:
        assert sheet.freeze_panes == "A2"
        assert sheet.auto_filter.ref is not None
        assert all(
            sheet.row_dimensions[row_index].height == 20
            for row_index in range(1, sheet.max_row + 1)
        )


def test_export_is_not_limited_to_frontend_page_size(client, session):
    for index in range(55):
        add_offer(session, f"offer-{index}", f"ENTREPRISE {index:02d}")
    session.commit()

    workbook = workbook_from(client.get("/api/v1/commercial-leads/export.xlsx"))

    assert workbook["Prospects"].max_row == 56


def test_recent_export_respects_window_and_kind(client, session, monkeypatch):
    import app.api.commercial_leads as commercial_leads_api

    class FixedDatetime(datetime):
        @classmethod
        def now(cls, tz=None):
            return NOW.astimezone(tz) if tz else NOW.replace(tzinfo=None)

    monkeypatch.setattr(commercial_leads_api, "datetime", FixedDatetime)
    add_offer(session, "known-old", "ENTREPRISE CONNUE", first_seen_hours_ago=24 * 30, title="Comptable")
    add_offer(session, "known-new", "ENTREPRISE CONNUE", first_seen_hours_ago=30, title="Commercial")
    add_offer(session, "new-company", "NOUVELLE ENTREPRISE", first_seen_hours_ago=2)
    session.commit()

    in_24h = workbook_from(client.get(
        "/api/v1/commercial-leads/recent/export.xlsx?window_hours=24&kind=all"
    ))
    new_offers_48h = workbook_from(client.get(
        "/api/v1/commercial-leads/recent/export.xlsx?window_hours=48&kind=new_offers"
    ))

    assert [row["Entreprise"] for row in sheet_records(in_24h["Prospects"])] == ["NOUVELLE ENTREPRISE"]
    assert [row["Entreprise"] for row in sheet_records(new_offers_48h["Prospects"])] == ["ENTREPRISE CONNUE"]
    assert sheet_records(in_24h["Prospects"])[0]["Type de nouveauté"] == "Nouvelle entreprise"
    assert sheet_records(new_offers_48h["Prospects"])[0]["Type de nouveauté"] == "Nouvelle offre"
    exported_offers = sheet_records(new_offers_48h["Offres"])
    assert sum(row["Nouvelle récemment"] == "Oui" for row in exported_offers) == 1
    assert sum(row["Nouvelle récemment"] == "Non" for row in exported_offers) == 1


def test_search_run_export_matches_current_results(client, session):
    add_offer(session, "search", "SEARCH COMPANY")
    add_contact_facts(session, "search company", "SEARCH COMPANY")
    runner = SearchRunOrchestrator(clock=lambda: NOW)
    run = runner.create(session, requested_actionable_leads=1, brave_hard_cap=0)
    completed = runner.resume(session, run.id)
    expected = runner.results(session, completed.id)

    workbook = workbook_from(client.get(f"/api/v1/search-runs/{completed.id}/export.xlsx"))

    assert [row["Entreprise"] for row in sheet_records(workbook["Prospects"])] == [
        lead.company_name for lead in expected
    ]


@pytest.mark.parametrize("dangerous", ("=1+1", "+1+1", "-1+1", "--cmd", "@SUM(A1:A2)"))
def test_formula_like_external_text_is_never_an_excel_formula(client, session, dangerous):
    add_offer(session, "formula-prefix", dangerous)
    session.commit()

    workbook = workbook_from(client.get(
        "/api/v1/commercial-leads/export.xlsx?include_excluded=true"
    ))
    company_cell = workbook["Prospects"].cell(2, 1)

    assert company_cell.data_type == "s"
    assert company_cell.value.endswith(dangerous)
