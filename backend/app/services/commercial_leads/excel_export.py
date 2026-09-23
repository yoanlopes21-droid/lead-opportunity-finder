"""Safe, in-memory XLSX export of already-composed commercial lead facts."""

from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime, timezone
from io import BytesIO
from typing import Iterable, Optional, Sequence
from urllib.parse import urlparse

from openpyxl import Workbook
from openpyxl.cell.cell import Cell
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.services.commercial_leads.service import CommercialLead


XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@dataclass(frozen=True)
class CommercialExcelItem:
    lead: CommercialLead
    latest_new_opportunity_at: Optional[datetime] = None
    is_new_company_in_window: Optional[bool] = None
    new_offer_ids_in_window: frozenset[str] = frozenset()


PROSPECT_HEADERS = (
    "Entreprise", "SIREN", "Site web", "Secteur", "Localisation principale",
    "Taille", "Score /100", "Catégorie de score", "Nombre d’offres actives",
    "Dernière nouveauté", "Type de nouveauté", "Offre publiée la plus récente", "Rôles recrutés",
    "Relation employeur", "Statut commercial", "Dernier contact", "Prochaine relance",
    "Contact prioritaire", "Fonction du contact", "Canal recommandé", "Email",
    "Téléphone", "URL professionnelle", "Confiance contact",
    "Pourquoi ce prospect est intéressant", "Warnings / limites",
    "Données manquantes utiles", "Sources / provenance", "URLs de preuve",
)

OFFER_HEADERS = (
    "Entreprise", "SIREN", "Intitulé", "Localisation", "Commune", "Contrat",
    "Salaire", "Date de publication", "Première détection locale", "Dernière observation",
    "Âge (jours)", "Nombre d’observations", "Nombre de sources", "Sources",
    "URLs de preuve", "Nouvelle récemment", "Identifiants source",
)

CONTACT_HEADERS = (
    "Entreprise", "SIREN", "Nom / personne", "Fonction", "Pertinence", "Email",
    "Téléphone", "URL professionnelle", "Type de canal", "Scope", "Confiance",
    "Vérification", "Source / provenance", "URL de preuve", "Contexte utile",
    "Incertitude / warning",
)

_EMPLOYER_LABELS = {
    "direct_employer": "Employeur direct",
    "intermediary_suspected": "Intermédiaire possible",
    "intermediary": "Intermédiaire",
}
_RELATIONSHIP_LABELS = {
    "contacted": "Contacté", "awaiting_reply": "En attente de réponse",
    "follow_up": "À relancer", "interested": "Intéressé",
    "meeting_scheduled": "Rendez-vous planifié", "proposal_sent": "Proposition envoyée",
    "client": "Client", "no_current_need": "Pas de besoin actuel", "refused": "Refusé",
    "wrong_contact": "Mauvais contact", "do_not_contact": "Ne plus contacter",
}
_CHANNEL_LABELS = {
    "direct_email": "Email direct", "functional_email": "Email recrutement",
    "direct_phone": "Téléphone direct", "service_phone": "Téléphone du service",
    "company_switchboard": "Standard entreprise", "official_contact_page": "Page contact officielle",
    "professional_url": "URL professionnelle", "none": "",
}
_ROLE_LABELS = {
    "hr": "Ressources humaines", "recruitment": "Recrutement", "director": "Direction",
    "manager": "Manager", "other": "Autre",
}
_SECTOR_LABELS = {
    "private": "Secteur privé", "public": "Secteur public",
    "nonprofit": "Association / organisme à but non lucratif", "unknown": "",
}
_DANGEROUS_PREFIXES = ("=", "+", "-", "@", "\t", "\r", "\n")
_HEADER_FILL = PatternFill("solid", fgColor="173D35")
_HEADER_FONT = Font(name="Arial", size=10, bold=True, color="FFFFFF")
_BODY_FONT = Font(name="Arial", size=10, color="172A28")
_THIN_BORDER = Border(bottom=Side(style="thin", color="DCE5DF"))
_DATE_FORMAT = "dd/mm/yyyy hh:mm"
_ROW_HEIGHT = 20


def build_commercial_xlsx(
    items: Sequence[CommercialExcelItem], *, generated_at: Optional[datetime] = None
) -> bytes:
    """Build a real XLSX in memory; callers own filtering and ordering."""
    generated_at = _as_utc(generated_at or datetime.now(timezone.utc))
    workbook = Workbook()
    prospects = workbook.active
    prospects.title = "Prospects"
    offers = workbook.create_sheet("Offres")
    contacts = workbook.create_sheet("Contacts")

    _write_sheet(prospects, PROSPECT_HEADERS, (_prospect_row(item) for item in items))
    _write_sheet(offers, OFFER_HEADERS, _offer_rows(items))
    _write_sheet(contacts, CONTACT_HEADERS, _contact_rows(items))

    workbook.properties.creator = "Lead Opportunity Finder"
    workbook.properties.title = "Export commercial"
    workbook.properties.created = generated_at.replace(tzinfo=None)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def export_filename(*, generated_at: Optional[datetime] = None, suffix: Optional[str] = None) -> str:
    date = _as_utc(generated_at or datetime.now(timezone.utc)).date().isoformat()
    middle = f"-{suffix}" if suffix else ""
    return f"lead-opportunity-finder{middle}-{date}.xlsx"


def _prospect_row(item: CommercialExcelItem) -> tuple[object, ...]:
    lead = item.lead
    strategy = lead.contact_strategy
    people = {person.id: person for person in lead.contactability.people}
    preferred_person = people.get(strategy.person_contact_id) if strategy else None
    usable_points = tuple(
        point for point in lead.contactability.contact_points
        if point.is_active and point.verification_status != "rejected"
    )
    preferred_point = next(
        (point for point in usable_points if strategy and point.id == strategy.contact_point_id), None
    )
    related_points = tuple(
        point for point in usable_points
        if preferred_person is not None and point.person_contact_id == preferred_person.id
    )
    candidate_points = _unique_points((*related_points, preferred_point, *usable_points))
    email = _first_point_value(candidate_points, "email")
    phone = _first_point_value(candidate_points, "phone")
    professional_url = _first_point_value(candidate_points, "professional_url")
    site = _official_site_url(lead) or _first_http_point(candidate_points, "website")
    relationship = lead.commercial_relationship
    reasons = _join_unique(
        item.message for item in lead.scoring.positive_reasons
    )
    signals = _join_unique(signal.explanation for signal in lead.latent_signals if signal.active)
    warnings = _join_unique((
        *(reason.message for reason in lead.scoring.penalties),
        *((strategy.warnings if strategy else ())),
        *((lead.exclusion.reason,) if lead.exclusion and lead.exclusion.reason else ()),
    ))
    sources = _join_unique((
        *(evidence.source_name for evidence in lead.evidence),
        *((reference.source_name for reference in strategy.evidence_references) if strategy else ()),
    ))
    proof_urls = _join_urls((
        *(evidence.source_url for evidence in lead.evidence),
        *((reference.source_url for reference in strategy.evidence_references) if strategy else ()),
        site,
    ))
    return (
        lead.company_name,
        lead.siren,
        site,
        _SECTOR_LABELS.get(lead.entity_sector_type, lead.entity_sector_type),
        lead.principal_location,
        None if lead.employee_range == "unknown" else lead.employee_range,
        lead.scoring.total_score,
        lead.scoring.category,
        lead.active_offer_count,
        item.latest_new_opportunity_at,
        (
            "Nouvelle entreprise" if item.is_new_company_in_window is True
            else "Nouvelle offre" if item.is_new_company_in_window is False
            and item.latest_new_opportunity_at is not None
            else None
        ),
        _parse_datetime(lead.newest_offer_created_at),
        _join_unique(lead.representative_job_titles),
        _EMPLOYER_LABELS.get(
            lead.scoring.employer_relationship_status,
            lead.scoring.employer_relationship_status,
        ),
        _RELATIONSHIP_LABELS.get(relationship.status, relationship.status) if relationship else None,
        relationship.last_contact_at if relationship else None,
        relationship.next_action_at if relationship else None,
        preferred_person.full_name if preferred_person else None,
        (preferred_person.job_title or _ROLE_LABELS.get(preferred_person.relevance_role)) if preferred_person else None,
        _CHANNEL_LABELS.get(strategy.preferred_channel, strategy.preferred_channel) if strategy else None,
        email,
        phone,
        professional_url,
        strategy.confidence if strategy else None,
        _join_unique((reasons, signals, strategy.short_context if strategy else None)),
        warnings,
        _join_unique(strategy.missing_information if strategy else ()),
        sources,
        proof_urls,
    )


def _offer_rows(items: Sequence[CommercialExcelItem]) -> Iterable[tuple[object, ...]]:
    for item in items:
        lead = item.lead
        for offer in lead.active_job_offers:
            technical_id = f"{offer.source}:{offer.offer_id}"
            is_new = "Oui" if technical_id in item.new_offer_ids_in_window else (
                "Non" if item.latest_new_opportunity_at is not None else None
            )
            yield (
                lead.company_name,
                lead.siren,
                offer.title,
                offer.display_location or offer.location_label,
                offer.commune,
                offer.contract_type,
                offer.salary,
                _parse_datetime(offer.published_at),
                offer.first_seen_at,
                offer.last_seen_at,
                offer.age_days,
                offer.observation_count,
                len(offer.sources),
                _join_unique(offer.sources),
                _join_urls(offer.source_urls),
                is_new,
                _join_unique(offer.source_offer_ids),
            )


def _contact_rows(items: Sequence[CommercialExcelItem]) -> Iterable[tuple[object, ...]]:
    for item in items:
        lead = item.lead
        people = {
            person.id: person for person in lead.contactability.people
            if person.is_active and person.verification_status != "rejected"
        }
        linked_person_ids: set[int] = set()
        for point in lead.contactability.contact_points:
            if not point.is_active or point.verification_status == "rejected":
                continue
            person = people.get(point.person_contact_id)
            if person:
                linked_person_ids.add(person.id)
            evidence = lead.contactability.evidence_by_contact_point_id.get(point.id, ())
            person_evidence = (
                lead.contactability.evidence_by_person_contact_id.get(person.id, ()) if person else ()
            )
            yield _contact_row(lead, person, point, (*person_evidence, *evidence))
        for person in people.values():
            if person.id not in linked_person_ids:
                evidence = lead.contactability.evidence_by_person_contact_id.get(person.id, ())
                yield _contact_row(lead, person, None, evidence)


def _contact_row(lead, person, point, evidence) -> tuple[object, ...]:
    point_type = point.contact_type if point else None
    value = point.normalized_value if point else None
    relevance = lead.contactability.channel_relevance_by_contact_point_id.get(point.id) if point else None
    warnings = []
    if person and person.verification_status in {"unverified", "stale"}:
        warnings.append("Personne ou fonction à vérifier.")
    if point and (point.verification_status == "stale" or point.confidence_level in {"review_needed", "ambiguous"}):
        warnings.append("Coordonnée à vérifier avant usage.")
    if relevance:
        warnings.extend(relevance.warnings)
    return (
        lead.company_name,
        lead.siren,
        person.full_name if person else None,
        person.job_title if person else None,
        _ROLE_LABELS.get(person.relevance_role, person.relevance_role) if person else None,
        value if point_type == "email" else None,
        value if point_type == "phone" else None,
        value if point_type in {"professional_url", "website"} else None,
        point_type,
        point.scope if point else person.scope if person else None,
        point.confidence_level if point else person.confidence_level if person else None,
        point.verification_status if point else person.verification_status if person else None,
        _join_unique((*((row.provider for row in evidence)), *((row.source_name for row in evidence)))),
        _join_urls(row.source_url for row in evidence),
        point.attribution_reason if point else person.attribution_reason if person else None,
        _join_unique(warnings),
    )


def _write_sheet(
    sheet: Worksheet, headers: Sequence[str], rows: Iterable[Sequence[object]]
) -> None:
    sheet.sheet_view.showGridLines = False
    sheet.append(tuple(headers))
    for row in rows:
        sheet.append(tuple(_excel_value(value) for value in row))
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(sheet.max_row, 1)}"
    for cell in sheet[1]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    date_columns = {
        index for index, header in enumerate(headers, start=1)
        if header in {
            "Dernière nouveauté", "Offre publiée la plus récente", "Dernier contact",
            "Prochaine relance", "Date de publication", "Première détection locale",
            "Dernière observation",
        }
    }
    hyperlink_columns = {
        index for index, header in enumerate(headers, start=1)
        if header in {"Site web", "URL professionnelle", "URL de preuve", "URLs de preuve"}
    }
    wrap_columns = {
        index for index, header in enumerate(headers, start=1)
        if header in {
            "Rôles recrutés", "Pourquoi ce prospect est intéressant", "Warnings / limites",
            "Données manquantes utiles", "Sources / provenance", "URLs de preuve",
            "Source / provenance", "URL de preuve", "Contexte utile", "Incertitude / warning",
        }
    }
    widths = [len(header) for header in headers]
    for row in sheet.iter_rows(min_row=2):
        for column, cell in enumerate(row, start=1):
            cell.font = _BODY_FONT
            cell.border = _THIN_BORDER
            cell.alignment = Alignment(
                vertical="top", wrap_text=(column in wrap_columns),
                horizontal="right" if isinstance(cell.value, (int, float)) else "left",
            )
            if column in date_columns and isinstance(cell.value, datetime):
                cell.number_format = _DATE_FORMAT
            if column in hyperlink_columns:
                _apply_hyperlink(cell)
            widths[column - 1] = min(max(widths[column - 1], _display_width(cell.value)), 48)
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = max(10, min(width + 2, 48))
    # Keep exports compact even when wrapped provenance or rationale is long.
    # This intentionally runs after every style/width adjustment.
    for row_index in range(1, sheet.max_row + 1):
        sheet.row_dimensions[row_index].height = _ROW_HEIGHT


def _excel_value(value: object) -> object:
    if isinstance(value, datetime):
        return _as_utc(value).replace(tzinfo=None)
    if isinstance(value, str):
        return _safe_text(value)
    return value


def _safe_text(value: str) -> str:
    # openpyxl writes +, -, @ and control-prefixed values as shared strings,
    # never formulas. A leading '=' is the one value it promotes to formula
    # XML, so prefix it explicitly. Whitespace-prefixed values remain strings.
    if not value.startswith(_DANGEROUS_PREFIXES):
        return value
    return f"'{value}" if value.startswith("=") else value


def _valid_http_url(value: Optional[str]) -> Optional[str]:
    if not isinstance(value, str) or not value.strip():
        return None
    cleaned = value.strip()
    parsed = urlparse(cleaned)
    return cleaned if parsed.scheme in {"http", "https"} and bool(parsed.netloc) else None


def _apply_hyperlink(cell: Cell) -> None:
    if not isinstance(cell.value, str) or "\n" in cell.value:
        return
    url = _valid_http_url(cell.value)
    if url:
        cell.hyperlink = url
        cell.style = "Hyperlink"


def _official_site_url(lead: CommercialLead) -> Optional[str]:
    candidates = [
        item for item in lead.contactability.verified_websites
        if item.status != "rejected" and _valid_http_url(item.canonical_url)
    ]
    if not candidates:
        return None
    priority = {"high_confidence": 0, "review_needed": 1, "ambiguous": 2}
    selected = min(candidates, key=lambda item: (priority.get(item.status, 99), -item.score, item.id))
    return _valid_http_url(selected.canonical_url)


def _first_point_value(points, contact_type: str) -> Optional[str]:
    point = next((item for item in points if item and item.contact_type == contact_type), None)
    return point.normalized_value if point else None


def _first_http_point(points, contact_type: str) -> Optional[str]:
    return next((url for item in points if item and item.contact_type == contact_type if (url := _valid_http_url(item.normalized_value))), None)


def _unique_points(points) -> tuple:
    return tuple({point.id: point for point in points if point is not None}.values())


def _join_unique(values: Iterable[Optional[str]]) -> Optional[str]:
    unique = tuple(dict.fromkeys(value.strip() for value in values if isinstance(value, str) and value.strip()))
    return "\n".join(unique) if unique else None


def _join_urls(values: Iterable[Optional[str]]) -> Optional[str]:
    return _join_unique(_valid_http_url(value) for value in values)


def _parse_datetime(value: Optional[str]) -> Optional[datetime]:
    if not value:
        return None
    try:
        parsed = datetime.fromisoformat(value.replace("Z", "+00:00"))
    except ValueError:
        return None
    return _as_utc(parsed)


def _as_utc(value: datetime) -> datetime:
    return value.astimezone(timezone.utc) if value.tzinfo else value.replace(tzinfo=timezone.utc)


def _display_width(value: object) -> int:
    if value is None:
        return 0
    return max((len(part) for part in str(value).splitlines()), default=0)
